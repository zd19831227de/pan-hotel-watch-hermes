from __future__ import annotations

import io
import unittest

from client.document_render import render_document, split_sections


_SAMPLE_MD = """# 静里酒店 10月房价铺排

## 一、分档结论
- 周中底价档：206 元（10-09 至 10-15）
- 周末上浮档：258-278 元（10-16、10-17 为本月最明确涨价窗口）

## 二、逐日铺排表
| 日期 | 建议价 | 依据 |
| 10-08 | 206-210 | 今日量弱，回落抢量 |
| 10-16 | 258-278 | 竞品集体 296，本店 226 偏低 |

## 三、风险与口径
调价须走审批流程，本表仅为建议区间。
"""


class DocumentRenderTests(unittest.TestCase):
    def test_split_sections_parses_headings_tables_bullets(self) -> None:
        sections = split_sections(_SAMPLE_MD)
        titles = [s["title"] for s in sections]
        self.assertEqual(titles[0], "静里酒店 10月房价铺排")
        self.assertTrue(any("逐日铺排表" in t for t in titles))
        table_section = next(s for s in sections if s.get("table"))
        self.assertEqual(table_section["table"][0], ["日期", "建议价", "依据"])
        self.assertEqual(len(table_section["table"]), 3)
        bullet_section = next(s for s in sections if s.get("bullets"))
        self.assertTrue(any("206" in b for b in bullet_section["bullets"]))

    def test_render_xlsx_returns_valid_workbook(self) -> None:
        payload = render_document(_SAMPLE_MD, doc_type="xlsx", title="静里铺排")
        self.assertIsInstance(payload, bytes)
        self.assertGreater(len(payload), 2000)
        import openpyxl

        book = openpyxl.load_workbook(io.BytesIO(payload))
        sheet_titles = book.sheetnames
        self.assertTrue(any("铺排" in name for name in sheet_titles), sheet_titles)
        # 表格行写入对应 sheet：表头 + 数据行。
        target = next(name for name in sheet_titles if "逐日" in name)
        rows = list(book[target].iter_rows(values_only=True))
        flat = [str(cell) for row in rows for cell in row if cell is not None]
        self.assertTrue(any("建议价" in cell for cell in flat), flat[:12])
        self.assertTrue(any("258-278" in cell for cell in flat))

    def test_render_docx_returns_valid_document(self) -> None:
        payload = render_document(_SAMPLE_MD, doc_type="docx", title="静里铺排")
        self.assertIsInstance(payload, bytes)
        self.assertGreater(len(payload), 2000)
        import docx

        document = docx.Document(io.BytesIO(payload))
        texts = [p.text for p in document.paragraphs if p.text.strip()]
        self.assertTrue(any("静里酒店" in t for t in texts), texts[:5])
        self.assertGreaterEqual(len(document.tables), 1)
        header = [cell.text for cell in document.tables[0].rows[0].cells]
        self.assertIn("建议价", header)

    def test_render_pptx_returns_valid_presentation(self) -> None:
        payload = render_document(_SAMPLE_MD, doc_type="pptx", title="静里铺排")
        self.assertIsInstance(payload, bytes)
        self.assertGreater(len(payload), 20000)
        from pptx import Presentation

        prs = Presentation(io.BytesIO(payload))
        self.assertGreaterEqual(len(prs.slides), 3)
        all_text = "\n".join(
            shape.text_frame.text
            for slide in prs.slides
            for shape in slide.shapes
            if shape.has_text_frame
        )
        self.assertIn("逐日铺排表", all_text)
        self.assertIn("206", all_text)

    def test_render_unknown_type_raises_value_error(self) -> None:
        with self.assertRaises(ValueError):
            render_document(_SAMPLE_MD, doc_type="exe", title="x")

    def test_render_pdf_returns_valid_pdf(self) -> None:
        # 2026-10-08 全链打通：PDF 输出（reportlab + 黑体/雅黑注册）。
        payload = render_document(_SAMPLE_MD, doc_type="pdf", title="静里铺排")
        self.assertIsInstance(payload, bytes)
        self.assertGreater(len(payload), 2000)
        self.assertTrue(payload.startswith(b"%PDF"))

    def test_pdf_safe_text_replaces_glyphs_missing_from_cjk_font(self) -> None:
        # 2026-10-09 生产质检：simhei 无 U+00A5（¥）/U+2022（•）字形，reportlab
        # 不做字体回退，直接渲染成方框。渲染前须归一为全角￥与中点·。
        from client.document_render import _pdf_safe_text

        self.assertEqual(_pdf_safe_text("¥265 • 要点"), "￥265 · 要点")
        self.assertEqual(_pdf_safe_text("普通中文 206-210 元"), "普通中文 206-210 元")

    def test_render_pdf_text_uses_cjk_safe_glyphs(self) -> None:
        md = "# 快报\n\n## 要点\n- 本店 ¥234 含双早\n\n## 表\n| 酒店 | 价 |\n| 立瑞 | ¥234 |\n"
        payload = render_document(md, doc_type="pdf", title="字形检查")
        try:
            import fitz  # noqa: F401
        except ImportError:
            self.skipTest("pymupdf 未安装（可选依赖）")
        doc = fitz.open(stream=payload, filetype="pdf")
        text = "".join(page.get_text() for page in doc)
        self.assertIn("￥234", text)
        self.assertNotIn("\u00a5", text)
        self.assertNotIn("\u2022", text)

    def test_render_pptx_table_section_uses_real_table(self) -> None:
        # 2026-10-09 生产质检：表格节必须是真表格（GraphicFrame），不再是
        # 「 | 」拼接的文本行。
        payload = render_document(_SAMPLE_MD, doc_type="pptx", title="静里铺排")
        from pptx import Presentation

        prs = Presentation(io.BytesIO(payload))
        tables = [
            shape.table
            for slide in prs.slides
            for shape in slide.shapes
            if shape.has_table
        ]
        self.assertEqual(len(tables), 1)
        table = tables[0]
        self.assertEqual(len(table.rows), 3)
        header = [cell.text for cell in table.rows[0].cells]
        self.assertEqual(header, ["日期", "建议价", "依据"])

    def test_render_drops_contentless_duplicate_title_section(self) -> None:
        # markdown 首个一级标题常与文档标题相同：不应产出只有标题的空节
        # （pptx 空页 / docx 标题重复）。
        title = "静里酒店 10月房价铺排"
        payload = render_document(_SAMPLE_MD, doc_type="docx", title=title)
        import docx

        document = docx.Document(io.BytesIO(payload))
        headings = [
            p.text
            for p in document.paragraphs
            if p.style.name.startswith(("Title", "Heading"))
        ]
        self.assertEqual(
            sum(1 for h in headings if h == title), 1, headings
        )

        pptx_payload = render_document(_SAMPLE_MD, doc_type="pptx", title=title)
        from pptx import Presentation

        prs = Presentation(io.BytesIO(pptx_payload))
        for index, slide in enumerate(prs.slides):
            if index == 0:
                continue  # 封面
            has_content = any(
                (shape.has_table and len(shape.table.rows) > 0)
                or (shape.has_text_frame and shape.text_frame.text.strip())
                for shape in slide.shapes
            )
            self.assertTrue(has_content, f"slide {index + 1} 为空页")

    def test_suggested_filename(self) -> None:
        from client.document_render import suggested_filename

        self.assertTrue(suggested_filename("静里铺排 10月", "xlsx").endswith(".xlsx"))
        self.assertIn("静里", suggested_filename("静里铺排", "pptx"))
        self.assertTrue(suggested_filename("国庆报告", "pdf").endswith(".pdf"))


if __name__ == "__main__":
    unittest.main()
