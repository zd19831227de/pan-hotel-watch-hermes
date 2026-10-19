"""
Turns the JSON produced by extract.py into a real, fully-editable .pptx
using native shapes (rectangles, text boxes, pictures, tables, connectors).

Usage:
    python build_pptx.py <slides.json> <output.pptx>
"""
import sys
import json
import os
import io
import base64
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

PX_TO_EMU = 9525  # 1 CSS px @ 96dpi -> EMU

ALIGN_MAP = {
    "left": PP_ALIGN.LEFT,
    "right": PP_ALIGN.RIGHT,
    "center": PP_ALIGN.CENTER,
    "justify": PP_ALIGN.JUSTIFY,
    "start": PP_ALIGN.LEFT,
    "end": PP_ALIGN.RIGHT,
}


def emu(px):
    return Emu(int(round(px * PX_TO_EMU)))


def hex_to_rgb(h):
    return RGBColor.from_string(h)


def set_alpha(color_format, alpha):
    """python-pptx has no public alpha/transparency setter; inject <a:alpha>."""
    if alpha >= 1.0:
        return
    srgb = color_format._xFill.find(qn("a:srgbClr"))
    if srgb is None:
        return
    for tag in ("a:alpha",):
        existing = srgb.find(qn(tag))
        if existing is not None:
            srgb.remove(existing)
    node = srgb.makeelement(qn("a:alpha"), {"val": str(int(alpha * 100000))})
    srgb.append(node)


def apply_fill(shape, el):
    if el.get("gradient"):
        g = el["gradient"]
        shape.fill.gradient()
        stops = shape.fill.gradient_stops
        stops[0].color.rgb = hex_to_rgb(g["from"])
        stops[-1].color.rgb = hex_to_rgb(g["to"])
        try:
            shape.fill.gradient_angle = g["angle"] % 360
        except Exception:
            pass
        return
    if el.get("fill"):
        shape.fill.solid()
        shape.fill.fore_color.rgb = hex_to_rgb(el["fill"])
        combined_alpha = el.get("fill_alpha", 1.0) * el.get("opacity", 1.0)
        set_alpha(shape.fill.fore_color, combined_alpha)
    else:
        shape.fill.background()


def apply_line(shape, el):
    if el.get("border_color") and el.get("border_width", 0) > 0:
        shape.line.color.rgb = hex_to_rgb(el["border_color"])
        shape.line.width = Pt(max(el["border_width"], 0.25) * 0.75)
    else:
        shape.line.fill.background()


def apply_shadow(shape, el):
    # get_or_add_effectLst() both suppresses theme-inherited effects (idx=2 in
    # <p:style> can carry a default shadow) and, unlike a raw append, reuses
    # the same element instead of risking a second sibling <a:effectLst>
    # (invalid per schema — CT_ShapeProperties allows at most one).
    spPr = shape._element.spPr
    effectLst = spPr.get_or_add_effectLst()
    if not el.get("shadow"):
        return
    outerShdw = effectLst.makeelement(qn("a:outerShdw"), {
        "blurRad": "90000", "dist": "30000", "dir": "5400000", "rotWithShape": "0",
    })
    clr = effectLst.makeelement(qn("a:srgbClr"), {"val": "000000"})
    alpha = clr.makeelement(qn("a:alpha"), {"val": "35000"})
    clr.append(alpha)
    outerShdw.append(clr)
    effectLst.append(outerShdw)


def add_rect(slide, el):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if el.get("radius", 0) > 2 else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        shape_type, emu(el["x"]), emu(el["y"]), emu(el["w"]), emu(el["h"])
    )
    if shape_type == MSO_SHAPE.ROUNDED_RECTANGLE and el["h"] > 0:
        try:
            shape.adjustments[0] = min(0.5, el["radius"] / min(el["w"], el["h"]))
        except Exception:
            pass
    apply_fill(shape, el)
    apply_line(shape, el)
    apply_shadow(shape, el)
    if el.get("rotation"):
        shape.rotation = el["rotation"]
    shape.text_frame.word_wrap = True
    return shape


def add_text(slide, el):
    tb = slide.shapes.add_textbox(emu(el["x"]), emu(el["y"]), emu(el["w"]), emu(el["h"]))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    paragraphs = el.get("paragraphs", [])
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN_MAP.get(para.get("align"), PP_ALIGN.LEFT)
        for run_data in para["runs"]:
            r = p.add_run()
            r.text = run_data["text"]
            f = r.font
            f.size = Pt(max(run_data.get("font_size", 12), 1))
            f.bold = run_data.get("bold", False)
            f.italic = run_data.get("italic", False)
            f.underline = run_data.get("underline", False)
            f.name = run_data.get("font_family", "Arial")
            f.color.rgb = hex_to_rgb(run_data.get("color", "000000"))
    if el.get("rotation"):
        tb.rotation = el["rotation"]
    return tb


def add_picture(slide, el, base_dir):
    src = el.get("src")
    if not src:
        return None
    image_source = None
    if src.startswith("data:"):
        header, _, b64data = src.partition(",")
        image_source = io.BytesIO(base64.b64decode(b64data))
    elif src.startswith(("http://", "https://")):
        print(f"WARNING: remote image URLs are not fetched, skipping: {src}", file=sys.stderr)
        return None
    else:
        path = src if os.path.isabs(src) else os.path.join(base_dir, src)
        if not os.path.exists(path):
            print(f"WARNING: image not found, skipping: {path}", file=sys.stderr)
            return None
        image_source = path
    pic = slide.shapes.add_picture(image_source, emu(el["x"]), emu(el["y"]), emu(el["w"]), emu(el["h"]))
    if el.get("rotation"):
        pic.rotation = el["rotation"]
    return pic


def add_line(slide, el):
    horizontal = el["w"] >= el["h"]
    x1, y1 = el["x"], el["y"]
    if horizontal:
        x2, y2 = el["x"] + el["w"], el["y"]
    else:
        x2, y2 = el["x"], el["y"] + el["h"]
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu(x1), emu(y1), emu(x2), emu(y2))
    color = el.get("fill") or el.get("border_color") or "000000"
    connector.line.color.rgb = hex_to_rgb(color)
    connector.line.width = Pt(max(el.get("border_width") or el.get("h") or 1, 0.75) * 0.75)
    return connector


def add_table(slide, el):
    rows_data = el.get("rows", [])
    if not rows_data:
        return None
    nrows = len(rows_data)
    ncols = max(len(r) for r in rows_data)
    graphic_frame = slide.shapes.add_table(
        nrows, ncols, emu(el["x"]), emu(el["y"]), emu(el["w"]), emu(el["h"])
    )
    table = graphic_frame.table
    for ri, row in enumerate(rows_data):
        for ci in range(ncols):
            cell = table.cell(ri, ci)
            if ci >= len(row):
                cell.text = ""
                continue
            cdata = row[ci]
            cell.text = cdata.get("text", "")
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            if cdata.get("fill"):
                cell.fill.solid()
                cell.fill.fore_color.rgb = hex_to_rgb(cdata["fill"])
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
            for p in cell.text_frame.paragraphs:
                p.alignment = ALIGN_MAP.get(cdata.get("align"), PP_ALIGN.LEFT)
                for r in p.runs:
                    r.font.bold = cdata.get("bold", False)
                    r.font.size = Pt(max(cdata.get("font_size", 12), 1))
                    r.font.color.rgb = hex_to_rgb(cdata.get("color", "000000"))
    return graphic_frame


def add_element(slide, el, base_dir):
    kind = el["kind"]
    if kind in ("rect", "background"):
        shape = add_rect(slide, el)
        if el.get("paragraphs"):
            # a background/rect that also carries text runs directly
            tf = shape.text_frame
            tf.word_wrap = True
            for i, para in enumerate(el["paragraphs"]):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.alignment = ALIGN_MAP.get(para.get("align"), PP_ALIGN.LEFT)
                for run_data in para["runs"]:
                    r = p.add_run()
                    r.text = run_data["text"]
                    r.font.size = Pt(max(run_data.get("font_size", 12), 1))
                    r.font.bold = run_data.get("bold", False)
                    r.font.color.rgb = hex_to_rgb(run_data.get("color", "000000"))
    elif kind == "text":
        add_text(slide, el)
    elif kind in ("image", "icon"):
        add_picture(slide, el, base_dir)
    elif kind == "line":
        add_line(slide, el)
    elif kind == "table":
        add_table(slide, el)
    else:
        print(f"WARNING: unknown data-pptx kind '{kind}', skipping", file=sys.stderr)


def main():
    json_path, out_path = sys.argv[1], sys.argv[2]
    base_dir = os.path.dirname(os.path.abspath(json_path))
    with open(json_path, encoding="utf-8") as f:
        slides_data = json.load(f)

    prs = Presentation()
    prs.slide_width = emu(slides_data[0]["width"])
    prs.slide_height = emu(slides_data[0]["height"])
    blank_layout = prs.slide_layouts[6]

    for slide_data in slides_data:
        slide = prs.slides.add_slide(blank_layout)
        for el in slide_data["elements"]:
            try:
                add_element(slide, el, base_dir)
            except Exception as exc:
                print(f"WARNING: failed to add element {el.get('kind')}: {exc}", file=sys.stderr)

    prs.save(out_path)
    print(f"Wrote {len(slides_data)} slide(s) -> {out_path}")


if __name__ == "__main__":
    main()
