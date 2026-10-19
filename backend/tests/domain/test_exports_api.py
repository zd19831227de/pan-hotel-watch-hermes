from __future__ import annotations

import importlib
import os
import unittest
from types import SimpleNamespace
from typing import Any
from unittest import mock

from fastapi.testclient import TestClient

_SUPERVISOR_KEY = "exports-supervisor-test-key"


class _SupervisorKeyEnv:
    def __enter__(self):
        self._saved = os.environ.get("PAN_HOTEL_BOOTSTRAP_SUPERVISOR_API_KEY")
        os.environ["PAN_HOTEL_BOOTSTRAP_SUPERVISOR_API_KEY"] = _SUPERVISOR_KEY
        settings = importlib.import_module("app.core.settings")
        settings.get_settings.cache_clear()
        return self

    def __exit__(self, *exc):
        if self._saved is None:
            os.environ.pop("PAN_HOTEL_BOOTSTRAP_SUPERVISOR_API_KEY", None)
        else:
            os.environ["PAN_HOTEL_BOOTSTRAP_SUPERVISOR_API_KEY"] = self._saved
        settings = importlib.import_module("app.core.settings")
        settings.get_settings.cache_clear()
        return False


class ExportsApiTests(unittest.TestCase):
    """UI-Fix-F2AC — 四个 CSV 导出端点 + 单次蒸馏导出。"""

    def _client(self) -> TestClient:
        main = importlib.import_module("app.main")
        return TestClient(main.create_app())

    def test_competitor_price_trend_csv(self) -> None:
        client = self._client()
        exports_api = importlib.import_module("app.api.v1.exports")
        trend: dict[str, Any] = {
            "status": "success",
            "watch_group_key": "default_competitor_watch",
            "granularity": "run",
            "points": [
                {
                    "bucket_at": "2026-08-20T10:00:00",
                    "hotel_name": "本店",
                    "hotel_role": "own",
                    "date_label": "today",
                    "checkin_date": "2026-08-20",
                    "breakfast_kind": "no_breakfast",
                    "room_type": "高级大床房",
                    "min_price_amount_yuan": 420.0,
                    "max_price_amount_yuan": 430.0,
                    "trusted_rate_count": 2,
                }
            ],
        }
        with mock.patch.object(
            exports_api.competitor_prices_repository,
            "get_competitor_price_trend",
            new=mock.AsyncMock(return_value=trend),
        ) as trend_mock:
            response = client.get(
                "/api/v1/export/competitor-price-trend.csv",
                params={"key": "default_competitor_watch", "days": 7},
            )
        self.assertEqual(response.status_code, 200)
        self.assertIn("text/csv", response.headers["content-type"])
        self.assertIn(
            "attachment; filename=", response.headers.get("content-disposition", "")
        )
        text = response.content.decode("utf-8-sig")
        self.assertIn("本店", text)
        self.assertIn("高级大床房", text)
        self.assertIn("420", text)
        kwargs = trend_mock.call_args.kwargs
        self.assertEqual(kwargs["watch_group_key"], "default_competitor_watch")
        self.assertEqual(kwargs["days"], 7)

    def test_price_change_alerts_csv(self) -> None:
        client = self._client()
        exports_api = importlib.import_module("app.api.v1.exports")
        alerts = [
            {
                "id": 1,
                "created_at": "2026-08-20T09:00:00",
                "watch_group_key": "g1",
                "status": "alert",
                "summary": "竞品降价 8%",
            }
        ]
        with mock.patch.object(
            exports_api.price_change_alerts_service,
            "list_price_change_alerts",
            new=mock.AsyncMock(return_value=alerts),
        ):
            response = client.get("/api/v1/export/price-change-alerts.csv?limit=50")
        self.assertEqual(response.status_code, 200)
        self.assertIn("text/csv", response.headers["content-type"])
        text = response.content.decode("utf-8-sig")
        self.assertIn("竞品降价 8%", text)

    def test_ctrip_ebk_runs_csv(self) -> None:
        client = self._client()
        exports_api = importlib.import_module("app.api.v1.exports")
        runs = {
            "watch_group_key": "ebk_g1",
            "runs": [
                {
                    "id": 11,
                    "captured_at": "2026-08-20T08:00:00",
                    "status": "success",
                    "summary": "9/9 datasets read",
                    "dataset_count": 9,
                    "success_count": 9,
                    "failed_count": 0,
                    "human_required_count": 0,
                    "source": "scheduler",
                    "actor_id": "scheduler",
                }
            ],
        }
        with mock.patch.object(
            exports_api.ebk_watch_repository,
            "list_data_center_watch_runs",
            new=mock.AsyncMock(return_value=runs),
        ):
            response = client.get("/api/v1/export/ctrip-ebk-runs.csv?key=ebk_g1&limit=20")
        self.assertEqual(response.status_code, 200)
        text = response.content.decode("utf-8-sig")
        self.assertIn("9/9 datasets read", text)
        self.assertIn("success", text)

    def test_distill_runs_csv(self) -> None:
        client = self._client()
        exports_api = importlib.import_module("app.api.v1.exports")
        distill_runs = [
            {
                "id": 3,
                "created_at": "2026-08-20T07:00:00",
                "platform": "doubao",
                "source_type": "text",
                "status": "success",
                "knowledge_title": "活动规则",
                "rag_document_id": 9,
                "reject_reason": None,
            }
        ]
        with mock.patch.object(
            exports_api.distill_repository,
            "list_distill_runs",
            new=mock.AsyncMock(return_value=distill_runs),
        ):
            response = client.get("/api/v1/export/distill-runs.csv?limit=20")
        self.assertEqual(response.status_code, 200)
        text = response.content.decode("utf-8-sig")
        self.assertIn("活动规则", text)
        self.assertIn("doubao", text)

    def test_exports_never_leak_sensitive_field_names(self) -> None:
        client = self._client()
        exports_api = importlib.import_module("app.api.v1.exports")
        with mock.patch.object(
            exports_api.competitor_prices_repository,
            "get_competitor_price_trend",
            new=mock.AsyncMock(return_value={"points": []}),
        ), mock.patch.object(
            exports_api.price_change_alerts_service,
            "list_price_change_alerts",
            new=mock.AsyncMock(return_value=[]),
        ), mock.patch.object(
            exports_api.ebk_watch_repository,
            "list_data_center_watch_runs",
            new=mock.AsyncMock(return_value={"runs": []}),
        ), mock.patch.object(
            exports_api.distill_repository,
            "list_distill_runs",
            new=mock.AsyncMock(return_value=[]),
        ):
            for path in (
                "/api/v1/export/competitor-price-trend.csv",
                "/api/v1/export/price-change-alerts.csv",
                "/api/v1/export/ctrip-ebk-runs.csv",
                "/api/v1/export/distill-runs.csv",
            ):
                response = client.get(path)
                self.assertEqual(response.status_code, 200, path)
                raw = response.content.decode("utf-8-sig").lower()
                for forbidden in (
                    "cookie",
                    "password",
                    "webhook",
                    "profile_path",
                    "browser_profile",
                    "api_key",
                    "secret",
                ):
                    self.assertNotIn(forbidden, raw, f"{path} leaked {forbidden}")

    def test_ctrip_ebk_runs_404_for_missing_group(self) -> None:
        client = self._client()
        exports_api = importlib.import_module("app.api.v1.exports")
        with mock.patch.object(
            exports_api.ebk_watch_repository,
            "list_data_center_watch_runs",
            new=mock.AsyncMock(return_value=None),
        ):
            response = client.get("/api/v1/export/ctrip-ebk-runs.csv?key=missing")
        self.assertEqual(response.status_code, 404)

    def test_price_change_alerts_csv_supports_xlsx_format(self) -> None:
        client = self._client()
        exports_api = importlib.import_module("app.api.v1.exports")
        with mock.patch.object(
            exports_api.price_change_alerts_service,
            "list_price_change_alerts",
            new=mock.AsyncMock(return_value=[]),
        ):
            response = client.get(
                "/api/v1/export/price-change-alerts.csv?format=xlsx"
            )
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            "spreadsheetml", response.headers["content-type"].lower()
        )
        self.assertTrue(
            response.headers.get("content-disposition", "").endswith('.xlsx"')
        )
        self.assertTrue(response.content.startswith(b"PK"))

    def test_ebk_business_data_export_csv_with_chinese_detail_rows(self) -> None:
        client = self._client()
        exports_api = importlib.import_module("app.api.v1.exports")
        from tests.domain.test_ebk_business_metrics import build_sample_snapshots

        with mock.patch.object(
            exports_api.ebk_snapshots_repository,
            "list_ebk_snapshots_with_payload_for_export",
            new=mock.AsyncMock(return_value=build_sample_snapshots()),
        ) as list_mock:
            response = client.get("/api/v1/export/ebk-business-data")
        self.assertEqual(response.status_code, 200)
        self.assertIn("text/csv", response.headers["content-type"])
        text = response.content.decode("utf-8-sig")
        self.assertIn("快照时间", text)
        self.assertIn("携程", text)
        self.assertIn("美团", text)
        self.assertIn("飞猪", text)
        self.assertIn("收益概览", text)
        self.assertIn("12400", text)
        self.assertNotIn("secret-screenshot-ref", text)
        kwargs = list_mock.call_args.kwargs
        self.assertEqual(kwargs["days"], 180)
        self.assertIsNone(kwargs["platforms"])

    def test_ebk_business_data_export_xlsx_dual_sheet_and_platform_filter(self) -> None:
        client = self._client()
        exports_api = importlib.import_module("app.api.v1.exports")
        from tests.domain.test_ebk_business_metrics import build_sample_snapshots

        snapshots = [
            snapshot
            for snapshot in build_sample_snapshots()
            if snapshot["platform"] != "ctrip"
        ]
        with mock.patch.object(
            exports_api.ebk_snapshots_repository,
            "list_ebk_snapshots_with_payload_for_export",
            new=mock.AsyncMock(return_value=snapshots),
        ) as list_mock:
            response = client.get(
                "/api/v1/export/ebk-business-data",
                params={"format": "xlsx", "days": 30, "platforms": "meituan,fliggy"},
            )
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response.headers["content-type"].lower())
        kwargs = list_mock.call_args.kwargs
        self.assertEqual(kwargs["days"], 30)
        self.assertEqual(kwargs["platforms"], ["meituan", "fliggy"])

        from openpyxl import load_workbook

        workbook = load_workbook(__import__("io").BytesIO(response.content))
        self.assertEqual(
            workbook.sheetnames,
            ["经营简报", "APP流量转化", "竞争圈排名", "核心指标对比", "经营明细", "指标说明"],
        )
        brief_sheet = workbook["经营简报"]
        self.assertEqual(brief_sheet["A1"].value, "EBK 经营数据简报")
        self.assertEqual(brief_sheet["A3"].value, "基础信息")
        self.assertEqual(brief_sheet["F3"].value, "数据简报")
        self.assertEqual(brief_sheet["L3"].value, "本月经营概览")
        self.assertEqual(brief_sheet["R3"].value, "竞争圈排名")
        self.assertEqual(brief_sheet["X3"].value, "APP流量转化")
        self.assertEqual(brief_sheet["A4"].value, "快照时间")
        self.assertEqual(brief_sheet["X4"].value, "列表页曝光量-我的酒店")
        self.assertIn("A3:E3", {str(item) for item in brief_sheet.merged_cells.ranges})
        self.assertIn("X3:AK3", {str(item) for item in brief_sheet.merged_cells.ranges})
        self.assertEqual(brief_sheet.freeze_panes, "A5")
        app_sheet = workbook["APP流量转化"]
        self.assertEqual(app_sheet["A3"].value, "基础信息")
        self.assertEqual(app_sheet["F3"].value, "APP流量转化")
        self.assertEqual(app_sheet["F4"].value, "列表页曝光量-我的酒店")
        self.assertEqual(app_sheet.freeze_panes, "A5")
        core_sheet = workbook["核心指标对比"]
        header_cells = [cell.value for cell in next(core_sheet.iter_rows(min_row=3, max_row=3))]
        self.assertEqual(header_cells[:3], ["快照时间", "平台", "账号"])
        platforms = [
            row[1] for row in core_sheet.iter_rows(min_row=4, values_only=True) if row[1]
        ]
        self.assertEqual(platforms, ["美团", "飞猪"])

    def test_ebk_business_brief_uses_configured_account_display_names(self) -> None:
        client = self._client()
        exports_api = importlib.import_module("app.api.v1.exports")
        from tests.domain.test_ebk_business_metrics import build_sample_snapshots

        snapshots = [
            snapshot
            for snapshot in build_sample_snapshots()
            if snapshot["platform"] in {"ctrip", "meituan"}
        ]
        sessions = [
            SimpleNamespace(
                platform="ctrip",
                platform_account_id="ctrip_main",
                display_name="静里携程ebk",
            ),
            SimpleNamespace(
                platform="meituan",
                platform_account_id="mt_main",
                display_name="静里美团ebk",
            ),
        ]
        with mock.patch.object(
            exports_api.ebk_snapshots_repository,
            "list_ebk_snapshots_with_payload_for_export",
            new=mock.AsyncMock(return_value=snapshots),
        ), mock.patch.object(
            exports_api,
            "get_browser_session_store",
            return_value=SimpleNamespace(list=lambda: sessions),
        ):
            response = client.get(
                "/api/v1/export/ebk-business-data",
                params={"format": "xlsx", "platforms": "ctrip,meituan"},
            )

        self.assertEqual(response.status_code, 200)
        from openpyxl import load_workbook

        workbook = load_workbook(__import__("io").BytesIO(response.content), read_only=True)
        brief_sheet = workbook["经营简报"]
        rows = list(brief_sheet.iter_rows(min_row=4, values_only=True))
        accounts = {row[2] for row in rows if row[2]}
        self.assertIn("静里携程ebk", accounts)
        self.assertIn("静里美团ebk", accounts)

    def test_pricing_strategy_history_export_flattens_suggestions(self) -> None:
        client = self._client()
        exports_api = importlib.import_module("app.api.v1.exports")
        alerts = [
            {
                "created_at": "2026-08-20T09:00:00",
                "watch_group_key": "g1",
                "suggestions": [
                    {
                        "hotel_name": "竞品A",
                        "room_type": "大床房",
                        "checkin_date": "2026-08-21",
                        "competitor_current_price_yuan": 380,
                        "current_own_price_yuan": 420,
                        "recommended_price_yuan": 400,
                        "recommended_adjustment_yuan": -20,
                        "trigger_reason": "竞品降价超过阈值",
                        "needs_human_attention": False,
                        "connection_mode": "non_direct",
                        "can_auto_write_price": True,
                    }
                ],
            },
            {"created_at": "2026-08-19T09:00:00", "watch_group_key": "g1", "suggestions": []},
        ]
        with mock.patch.object(
            exports_api.price_change_alerts_service,
            "list_price_change_alerts",
            new=mock.AsyncMock(return_value=alerts),
        ) as list_mock:
            response = client.get("/api/v1/export/pricing-strategy-history")
        self.assertEqual(response.status_code, 200)
        text = response.content.decode("utf-8-sig")
        self.assertIn("竞品A", text)
        self.assertIn("大床房", text)
        self.assertIn("400", text)
        self.assertNotIn("2026-08-19", text, "alerts without suggestions must not emit rows")
        kwargs = list_mock.call_args.kwargs
        self.assertEqual(kwargs["days"], 180)

    def test_new_exports_do_not_leak_sensitive_fields(self) -> None:
        client = self._client()
        exports_api = importlib.import_module("app.api.v1.exports")
        with mock.patch.object(
            exports_api.ebk_snapshots_repository,
            "list_ebk_snapshots_with_payload_for_export",
            new=mock.AsyncMock(return_value=[]),
        ), mock.patch.object(
            exports_api.price_change_alerts_service,
            "list_price_change_alerts",
            new=mock.AsyncMock(return_value=[]),
        ):
            for path in (
                "/api/v1/export/ebk-business-data",
                "/api/v1/export/pricing-strategy-history",
            ):
                response = client.get(path)
                self.assertEqual(response.status_code, 200, path)
                raw = response.content.decode("utf-8-sig", errors="ignore").lower()
                for forbidden in (
                    "cookie",
                    "password",
                    "webhook",
                    "profile_path",
                    "browser_profile",
                    "api_key",
                    "secret",
                    "browser_session_id",
                ):
                    self.assertNotIn(forbidden, raw, f"{path} leaked {forbidden}")


class DistillRunExportApiTests(unittest.TestCase):
    def _client(self) -> TestClient:
        main = importlib.import_module("app.main")
        return TestClient(main.create_app())

    def test_distill_run_export_requires_supervisor_key(self) -> None:
        client = self._client()
        response = client.get("/api/v1/distill/runs/1/export")
        self.assertIn(response.status_code, (401, 403))

    def test_distill_run_export_returns_knowledge_csv(self) -> None:
        run = {
            "id": 5,
            "created_at": "2026-08-20T06:00:00",
            "platform": "qwen",
            "source_type": "text",
            "status": "success",
            "summary": "蒸馏完成并入库。",
            "extracted_text": "1. 活动报名需提前 3 天\n2. 退款按平台规则",
            "knowledge_title": "活动规则",
            "rag_document_id": 12,
            "reject_reason": None,
        }
        distill_api = importlib.import_module("app.api.v1.distill")
        with _SupervisorKeyEnv():
            client = self._client()
            with mock.patch.object(
                distill_api.repository,
                "get_distill_run",
                new=mock.AsyncMock(return_value=run),
            ):
                response = client.get(
                    "/api/v1/distill/runs/5/export",
                    headers={"X-Pan-Hotel-API-Key": _SUPERVISOR_KEY},
                )
        self.assertEqual(response.status_code, 200)
        self.assertIn("text/csv", response.headers["content-type"])
        self.assertTrue(response.content.startswith(b"\xef\xbb\xbf"))
        text = response.content.decode("utf-8-sig")
        self.assertIn("活动规则", text)
        self.assertIn("提前 3 天", text)

    def test_distill_run_export_404_for_missing_run(self) -> None:
        distill_api = importlib.import_module("app.api.v1.distill")
        with _SupervisorKeyEnv():
            client = self._client()
            with mock.patch.object(
                distill_api.repository,
                "get_distill_run",
                new=mock.AsyncMock(return_value=None),
            ):
                response = client.get(
                    "/api/v1/distill/runs/99/export",
                    headers={"X-Pan-Hotel-API-Key": _SUPERVISOR_KEY},
                )
        self.assertEqual(response.status_code, 404)


if __name__ == "__main__":
    unittest.main()
