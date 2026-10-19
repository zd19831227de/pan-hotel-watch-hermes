from __future__ import annotations

import csv
import io
import unittest


class CsvBuilderTests(unittest.TestCase):
    """UI-Fix-F2AC — 通用 CSV 构造器：utf-8-sig BOM + RFC 4180 转义。"""

    def _load(self):
        from app.domain.exports.csv_builder import build_csv_bytes

        return build_csv_bytes

    def test_module_exists_and_outputs_bom(self) -> None:
        build_csv_bytes = self._load()
        payload = build_csv_bytes(headers=["id", "名称"], rows=[[1, "酒店A"]])
        self.assertTrue(payload.startswith(b"\xef\xbb\xbf"), "utf-8-sig BOM must exist")

    def test_chinese_fields_round_trip(self) -> None:
        build_csv_bytes = self._load()
        payload = build_csv_bytes(
            headers=["时间", "平台", "状态"],
            rows=[["2026-08-20 10:00:00", "携程", "成功"]],
        )
        text = payload.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
        self.assertEqual(rows[0], ["时间", "平台", "状态"])
        self.assertEqual(rows[1], ["2026-08-20 10:00:00", "携程", "成功"])

    def test_comma_quote_newline_escaping(self) -> None:
        build_csv_bytes = self._load()
        payload = build_csv_bytes(
            headers=["摘要"],
            rows=[["含,逗号"], ['含"双引号'], ["含\n换行"]],
        )
        text = payload.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
        self.assertEqual(rows[0], ["摘要"])
        self.assertEqual(rows[1], ["含,逗号"])
        self.assertEqual(rows[2], ['含"双引号'])
        self.assertEqual(rows[3], ["含\n换行"])
        self.assertIn('"含,逗号"', text)
        self.assertIn('"含""双引号"', text)

    def test_none_values_become_empty(self) -> None:
        build_csv_bytes = self._load()
        payload = build_csv_bytes(headers=["a", "b"], rows=[[None, 2]])
        rows = list(csv.reader(io.StringIO(payload.decode("utf-8-sig"))))
        self.assertEqual(rows[1], ["", "2"])

    def test_no_rows_still_has_header(self) -> None:
        build_csv_bytes = self._load()
        payload = build_csv_bytes(headers=["id"], rows=[])
        text = payload.decode("utf-8-sig").strip()
        self.assertEqual(text, "id")


class XlsxBuilderTests(unittest.TestCase):
    """Excel 导出（修正4）：openpyxl 构造，双击 Excel 直开。"""

    def _load(self):
        from app.domain.exports.xlsx_builder import build_xlsx_bytes

        return build_xlsx_bytes

    def test_xlsx_is_zip_container_with_headers_and_rows(self) -> None:
        build_xlsx_bytes = self._load()
        payload = build_xlsx_bytes(
            headers=["时间", "平台", "价格"],
            rows=[["2026-08-20", "携程", 420]],
        )
        self.assertTrue(payload.startswith(b"PK"), "xlsx must be a zip container")
        import io

        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(payload))
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(rows[0], ("时间", "平台", "价格"))
        self.assertEqual(rows[1], ("2026-08-20", "携程", 420))

    def test_xlsx_handles_none_and_long_text(self) -> None:
        build_xlsx_bytes = self._load()
        payload = build_xlsx_bytes(
            headers=["a", "b"],
            rows=[[None, "含,逗号\n换行\"引号"]],
        )
        import io

        from openpyxl import load_workbook

        sheet = load_workbook(io.BytesIO(payload)).active
        rows = list(sheet.iter_rows(values_only=True))
        self.assertIsNone(rows[1][0])
        self.assertEqual(rows[1][1], '含,逗号\n换行"引号')

    def test_xlsx_empty_rows_keeps_header(self) -> None:
        build_xlsx_bytes = self._load()
        payload = build_xlsx_bytes(headers=["id"], rows=[])
        import io

        from openpyxl import load_workbook

        sheet = load_workbook(io.BytesIO(payload)).active
        rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(rows, [("id",)])

    def test_xlsx_sheet_spec_supports_mobile_friendly_brief_style(self) -> None:
        from app.domain.exports.xlsx_builder import build_xlsx_workbook_bytes
        from openpyxl import load_workbook

        payload = build_xlsx_workbook_bytes(
            sheets=[
                {
                    "name": "经营简报",
                    "title": "EBK 经营数据简报",
                    "subtitle": "空白字段表示本次成功读取未提供该数据",
                    "headers": ["平台", "城市排名", "本月销售额"],
                    "rows": [{"平台": "携程", "城市排名": 280, "本月销售额": 562679.6}],
                    "freeze_panes": "A4",
                    "auto_filter": True,
                    "column_widths": {"平台": 12, "本月销售额": 18},
                }
            ]
        )
        sheet = load_workbook(io.BytesIO(payload))["经营简报"]
        self.assertEqual(sheet["A1"].value, "EBK 经营数据简报")
        self.assertEqual(sheet["A3"].value, "平台")
        self.assertEqual(sheet.freeze_panes, "A4")
        self.assertEqual(sheet.auto_filter.ref, "A3:C4")
        self.assertEqual(sheet.column_dimensions["A"].width, 12)
        self.assertTrue(sheet["A3"].font.bold)
        self.assertEqual(sheet["A3"].font.color.rgb, "00FFFFFF")
        self.assertNotEqual(sheet["A4"].fill.fill_type, None)


if __name__ == "__main__":
    unittest.main()
