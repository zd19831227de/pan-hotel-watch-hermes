"""通用 Excel（xlsx）构造器（修正4）：openpyxl 内存构造，双击 Excel 直开。"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.domain.exports.csv_builder import _stringify


def _cell_value(value: Any) -> Any:
    """xlsx 单元格保留数值类型（Excel 可直接计算），其余转字符串。"""
    if value is None:
        return None
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, float)):
        return value
    return _stringify(value)


def build_xlsx_bytes(
    *,
    headers: list[str],
    rows: list[list[Any]] | list[dict[str, Any]],
) -> bytes:
    """构造 xlsx 字节流；rows 语义与 build_csv_bytes 一致（list 或 dict 按 headers 取值）。"""
    return build_xlsx_workbook_bytes(
        sheets=[{"name": "export", "headers": headers, "rows": rows}]
    )


def build_xlsx_workbook_bytes(*, sheets: list[dict[str, Any]]) -> bytes:
    """构造多 Sheet xlsx：sheets = [{name, headers, rows}]，单元格语义同单表版本。"""

    workbook = Workbook()
    workbook.remove(workbook.active)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="245B9E")
    alternate_fill = PatternFill("solid", fgColor="F3F6FA")
    for sheet_spec in sheets or [{"name": "export", "headers": [], "rows": []}]:
        sheet = workbook.create_sheet(title=str(sheet_spec.get("name") or "export")[:31])
        headers = sheet_spec.get("headers") or []
        rows = sheet_spec.get("rows") or []
        title = sheet_spec.get("title")
        subtitle = sheet_spec.get("subtitle")
        header_row = 1
        if title:
            sheet.cell(row=1, column=1, value=_stringify(title)).font = Font(
                bold=True, size=16, color="17365D"
            )
            sheet.merge_cells(
                start_row=1,
                start_column=1,
                end_row=1,
                end_column=max(1, len(headers)),
            )
            header_row += 1
        if subtitle:
            sheet.cell(row=2, column=1, value=_stringify(subtitle)).font = Font(
                italic=True, color="667085"
            )
            sheet.merge_cells(
                start_row=2,
                start_column=1,
                end_row=2,
                end_column=max(1, len(headers)),
            )
            header_row += 1
        column_groups = sheet_spec.get("column_groups") or []
        if column_groups and headers:
            group_row = header_row
            header_positions = {
                str(header): column for column, header in enumerate(headers, start=1)
            }
            for group in column_groups:
                if not isinstance(group, dict):
                    continue
                start_column = header_positions.get(str(group.get("start") or ""))
                end_column = header_positions.get(str(group.get("end") or ""))
                if not start_column or not end_column or end_column < start_column:
                    continue
                if end_column > start_column:
                    sheet.merge_cells(
                        start_row=group_row,
                        start_column=start_column,
                        end_row=group_row,
                        end_column=end_column,
                    )
                fill = PatternFill(
                    "solid",
                    fgColor=str(group.get("fill") or "455A64").lstrip("#"),
                )
                for column in range(start_column, end_column + 1):
                    cell = sheet.cell(row=group_row, column=column)
                    cell.fill = fill
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                group_cell = sheet.cell(
                    row=group_row,
                    column=start_column,
                    value=_stringify(group.get("title") or ""),
                )
                group_cell.font = Font(bold=True, color="FFFFFF")
            sheet.row_dimensions[group_row].height = 24
            header_row += 1
        header_fills = sheet_spec.get("header_fills") or {}
        for column, header in enumerate(headers, start=1):
            cell = sheet.cell(row=header_row, column=column, value=_stringify(header))
            cell.font = header_font
            override = header_fills.get(header)
            cell.fill = (
                PatternFill("solid", fgColor=str(override).lstrip("#"))
                if override
                else header_fill
            )
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row_index, row in enumerate(rows, start=header_row + 1):
            if isinstance(row, dict):
                values = [_cell_value(row.get(header)) for header in headers]
            else:
                values = [_cell_value(cell) for cell in row]
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row=row_index, column=column, value=value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row_index % 2 == 0:
                    cell.fill = alternate_fill
        if headers:
            last_row = max(header_row, header_row + len(rows))
            if sheet_spec.get("auto_filter", True):
                sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{last_row}"
            sheet.freeze_panes = str(sheet_spec.get("freeze_panes") or f"A{header_row + 1}")
            widths = sheet_spec.get("column_widths") or {}
            for column, header in enumerate(headers, start=1):
                width = widths.get(header)
                if width is None:
                    samples = [str(header)] + [
                        str(row.get(header, ""))
                        if isinstance(row, dict)
                        else str(row[column - 1] if len(row) >= column else "")
                        for row in rows[:30]
                    ]
                    width = min(32, max(10, max(len(sample) for sample in samples) + 2))
                sheet.column_dimensions[get_column_letter(column)].width = float(width)
            sheet.row_dimensions[header_row].height = 30
        sheet.sheet_view.showGridLines = False
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
