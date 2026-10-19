# -*- coding: utf-8 -*-
"""方案A（2026-10-09）：高档渲染引擎（收编 pptx-builder/pdf-builder 管线）。

三层验证：①markdown→HTML 模板结构（data-pptx/字体/转义）；②开关与回退
（PAN_HOTEL_DOC_PRO_DISABLED=1 时回退轻量渲染器）；③端到端编译（可选，
依赖 playwright+chromium，缺失时跳过——生产 hermes venv 需装齐后生效）。
"""

from __future__ import annotations

import io
import os
import unittest
from unittest import mock

from client.document_render import render_document, split_sections
from client.render_pro import pro_disabled, render_document_pro
from client.render_pro.markdown_to_html import build_html

_MD = """# 静里酒店 10月房价铺排

## 一、分档结论
- 周中底价档：206 元（10-09 至 10-15）
- 周末上浮档：258-278 元（10-16、10-17 为本月最明确涨价窗口）

## 二、逐日铺排表
| 日期 | 建议价 | 依据 |
| 10-08 | 206-210 | 今日量弱，回落抢量 |
| 10-16 | 258-278 | 竞品集体 296，本店 226 偏低 |

## 三、风险与口径
调价须走审批流程，本表仅为建议区间。
"""


def _sections() -> list[dict]:
    return [s for s in split_sections(_MD) if s.get("table") or s.get("bullets") or s.get("paragraph")]


class MarkdownToHtmlTests(unittest.TestCase):
    def test_pptx_html_structure(self) -> None:
        html = build_html(_sections(), doc_type="pptx", title="静里铺排", meta_lines=["口径：竞品落库"])
        self.assertIn('class="slide"', html)
        self.assertIn('width:1280px;height:720px', html)
        self.assertIn('data-pptx="table"', html)
        self.assertIn('data-pptx="background"', html)
        self.assertIn("Microsoft YaHei", html)
        self.assertIn("静里铺排", html)
        # 无内容标题节不出封面外多余版面：封面 + 3 个内容节。
        self.assertEqual(html.count('class="slide"'), 4)

    def test_pdf_html_uses_simhei_and_data_pdf_attr(self) -> None:
        html = build_html(_sections(), doc_type="pdf", title="静里铺排")
        self.assertIn('class="page"', html)
        self.assertIn('width:816px;height:1056px', html)
        self.assertIn("'SimHei'", html)
        self.assertNotIn("Microsoft YaHei", html)
        # pdf 管线元素标记是 data-pdf（extract_pdf.py 的契约）。
        self.assertIn('data-pdf="table"', html)
        self.assertNotIn('data-pptx=', html)

    def test_html_escapes_user_text(self) -> None:
        sections = [{"title": "节", "bullets": ["<b>加粗注入</b> & 符号"], "paragraph": ""}]
        html = build_html(sections, doc_type="pptx", title="t")
        self.assertNotIn("<b>加粗注入</b>", html)
        self.assertIn("&lt;b&gt;加粗注入&lt;/b&gt;", html)


class ProEngineGateTests(unittest.TestCase):
    def test_disabled_switch_returns_none(self) -> None:
        with mock.patch.dict(os.environ, {"PAN_HOTEL_DOC_PRO_DISABLED": "1"}):
            self.assertTrue(pro_disabled())
            self.assertIsNone(render_document_pro(_MD, doc_type="pptx", title="t"))

    def test_unsupported_types_return_none(self) -> None:
        self.assertIsNone(render_document_pro(_MD, doc_type="xlsx", title="t"))
        self.assertIsNone(render_document_pro(_MD, doc_type="docx", title="t"))

    def test_render_document_falls_back_to_lightweight_when_disabled(self) -> None:
        # 回退路径 = 轻量渲染器：python-pptx 默认画布 10in（9144000 EMU）。
        with mock.patch.dict(os.environ, {"PAN_HOTEL_DOC_PRO_DISABLED": "1"}):
            payload = render_document(_MD, doc_type="pptx", title="静里铺排")
        from pptx import Presentation

        prs = Presentation(io.BytesIO(payload))
        self.assertEqual(prs.slide_width, 9144000)

    def test_render_document_falls_back_when_pro_raises(self) -> None:
        with mock.patch("client.render_pro.render_document_pro", side_effect=RuntimeError("boom")):
            # document_render 内 import 后调用 —— 打补丁在模块属性上生效。
            payload = render_document(_MD, doc_type="pdf", title="静里铺排")
        self.assertTrue(payload.startswith(b"%PDF"))


@unittest.skipUnless(
    __import__("importlib.util", fromlist=["util"]).find_spec("playwright") is not None,
    "playwright 未安装（可选依赖，生产 hermes venv 需装齐）",
)
class ProEngineE2ETests(unittest.TestCase):
    """端到端：markdown → HTML → Playwright 测量 → 原生文件。"""

    def test_pptx_pro_builds_native_wide_canvas_deck(self) -> None:
        payload = render_document(_MD, doc_type="pptx", title="静里酒店 10月房价铺排")
        from pptx import Presentation

        prs = Presentation(io.BytesIO(payload))
        # pro 引擎 16:9 画布：1280px@96dpi = 12192000 EMU（轻量版是 9144000）。
        self.assertEqual(prs.slide_width, 12192000)
        tables = [
            shape.table
            for slide in prs.slides
            for shape in slide.shapes
            if shape.has_table
        ]
        self.assertEqual(len(tables), 1)
        header = [cell.text for cell in tables[0].rows[0].cells]
        self.assertEqual(header, ["日期", "建议价", "依据"])

    def test_pdf_pro_builds_vector_pdf_with_cjk(self) -> None:
        payload = render_document(_MD, doc_type="pdf", title="静里酒店 10月房价铺排")
        self.assertTrue(payload.startswith(b"%PDF"))
        try:
            import fitz
        except ImportError:
            self.skipTest("pymupdf 未安装")
        doc = fitz.open(stream=payload, filetype="pdf")
        text = "".join(page.get_text() for page in doc)
        self.assertIn("静里", text)
        self.assertIn("258-278", text)
        self.assertNotIn("\x00", text)


if __name__ == "__main__":
    unittest.main()
